from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd
import pdfplumber
import re


@dataclass(frozen=True)
class InputBundle:
    plan_type: str
    folder: Path
    csv_file: Path
    pdf_files: list[Path]
    excel_files: list[Path]


_PLAN_TYPE_RE = re.compile(r"^(PP|PC)", flags=re.IGNORECASE)


def _detect_plan_type(file_path: Path) -> str | None:
    match = _PLAN_TYPE_RE.match(file_path.stem)
    return match.group(1).upper() if match else None


def locate_input_bundle(input_dir: Path, csv_pattern: str, pdf_pattern: str, excel_pattern: str) -> InputBundle:
    if not input_dir.exists():
        raise ValueError(f"Pasta de entrada nao existe: {input_dir}")

    csv_files = sorted(input_dir.glob(csv_pattern))
    if not csv_files:
        raise ValueError(f"Nenhum CSV encontrado em {input_dir}")
    if len(csv_files) > 1:
        raise ValueError(f"Mais de um CSV encontrado em {input_dir}: {[file.name for file in csv_files]}")

    pdf_files = sorted(input_dir.glob(pdf_pattern))
    excel_files = sorted(input_dir.glob(excel_pattern))

    plan_type = _detect_plan_type(csv_files[0])
    if plan_type is None:
        for candidate in [*pdf_files, *excel_files]:
            plan_type = _detect_plan_type(candidate)
            if plan_type:
                break
    if plan_type is None:
        raise ValueError(
            f"Nao foi possivel detectar tipo do plano: nenhum arquivo PP/PC encontrado "
            f"(csv={csv_files[0].name})"
        )

    return InputBundle(
        plan_type=plan_type,
        folder=input_dir,
        csv_file=csv_files[0],
        pdf_files=pdf_files,
        excel_files=excel_files,
    )


def read_csv(csv_file: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(csv_file, skipinitialspace=True)
    except Exception as exc:
        raise ValueError(f"Falha ao ler CSV {csv_file}: {exc}") from exc


def extract_plan_id_from_pdf(pdf_file: Path, pdf_regex: str) -> str | None:
    try:
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                plan_id = extract_plan_id_from_text(text, pdf_regex)
                if plan_id:
                    return plan_id
    except Exception as exc:
        raise ValueError(f"Falha ao ler PDF {pdf_file}: {exc}") from exc
    return None


def extract_plan_id_from_text(text: str, regex: str) -> str | None:
    match = re.search(regex, text, flags=re.IGNORECASE)
    if match:
        return match.group(1) if match.lastindex else match.group(0)
    return None


def extract_plan_id_from_filename(file_path: Path, filename_regex: str) -> str | None:
    return extract_plan_id_from_text(file_path.stem, filename_regex)


def extract_plan_id_from_excel(excel_file: Path, plan_regex: str) -> str | None:
    pattern = re.compile(plan_regex, flags=re.IGNORECASE)
    try:
        workbook = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    for value in row:
                        if value is None:
                            continue
                        match = pattern.search(str(value))
                        if match:
                            return match.group(1) if match.lastindex else match.group(0)
        finally:
            workbook.close()
    except Exception as exc:
        raise ValueError(f"Falha ao ler XLSX {excel_file}: {exc}") from exc
    return None


def extract_excel_column_map(excel_file: Path, key_column: str, value_column: str) -> dict[str, object]:
    try:
        workbook = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue
                header_index = {str(header): idx for idx, header in enumerate(headers) if header is not None}
                if key_column not in header_index or value_column not in header_index:
                    continue

                key_idx = header_index[key_column]
                value_idx = header_index[value_column]
                values: dict[str, object] = {}
                for row in rows:
                    if len(row) <= max(key_idx, value_idx):
                        continue
                    key = row[key_idx]
                    value = row[value_idx]
                    if key is None or value is None:
                        continue
                    values[str(key)] = value
                return values
        finally:
            workbook.close()
    except Exception as exc:
        raise ValueError(f"Falha ao ler XLSX {excel_file}: {exc}") from exc
    raise ValueError(f"Colunas obrigatorias nao encontradas em {excel_file}: {key_column}, {value_column}")
