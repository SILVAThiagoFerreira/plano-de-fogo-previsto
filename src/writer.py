from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd


def write_excel(df: pd.DataFrame, output_file: Path, sheet_name: str, column_width: int) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.sheets[sheet_name]
        for idx, _column in enumerate(df.columns, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = column_width


def apply_numeric_format(output_file: Path, sheet_name: str, precision: int, column_precisions: dict[str, int] | None = None) -> None:
    workbook = openpyxl.load_workbook(output_file)
    worksheet = workbook[sheet_name]
    default_format = "0" if precision <= 0 else "0." + ("0" * precision)
    column_precisions = column_precisions or {}
    header_to_format: dict[str, str] = {}
    for cell in worksheet[1]:
        if cell.value is None:
            continue
        header = str(cell.value)
        header_precision = column_precisions.get(header, precision)
        header_to_format[header] = "0" if header_precision <= 0 else "0." + ("0" * header_precision)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                header = worksheet.cell(row=1, column=cell.column).value
                cell.number_format = header_to_format.get(str(header), default_format)
    workbook.save(output_file)
